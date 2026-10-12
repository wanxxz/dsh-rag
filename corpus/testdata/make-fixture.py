"""生成中文 DOCX 测试样本（供 ingest.py 的 DOCX 路径回归用）。

用 DSH 自带 Python 运行：
  <bundled-python> make-fixture.py
"""
import os

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))

doc = Document()

# 关键：显式指定中文字体。默认字体（Calibri 等）没有中文字形，
# 导出 PDF 时会发生字体替换、把字序打乱，导致文本层不可用。
for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
    try:
        st = doc.styles[style_name]
        st.font.name = "SimSun"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    except KeyError:
        pass
doc.add_heading("关于开展新能源发电项目并网管理的通知", level=0)

doc.add_heading("第一章 总则", level=1)
doc.add_paragraph("第一条 为规范新能源发电项目并网管理，促进新能源高质量发展，制定本通知。")
doc.add_paragraph("第二条 本通知适用于接入配电网运行的分布式光伏、分散式风电项目。")

doc.add_heading("第二章 并网服务", level=1)
doc.add_paragraph("第三条 电网企业应当为新能源项目提供便捷高效的并网服务，不得无故拒绝或者拖延。")

doc.add_heading("一、办理时限", level=2)
doc.add_paragraph("第四条 电网企业应当自受理并网申请之日起二十个工作日内完成答复。")

table = doc.add_table(rows=3, cols=3)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = "项目类型", "容量上限", "办理时限"
r1 = table.rows[1].cells
r1[0].text, r1[1].text, r1[2].text = "分布式光伏", "6兆瓦", "20个工作日"
r2 = table.rows[2].cells
r2[0].text, r2[1].text, r2[2].text = "分散式风电", "50兆瓦", "30个工作日"

doc.add_heading("第三章 附则", level=1)
doc.add_paragraph("第五条 本通知自发布之日起施行，有效期五年。")

out = os.path.join(HERE, "sample-policy.docx")
doc.save(out)
print("已生成:", out)
